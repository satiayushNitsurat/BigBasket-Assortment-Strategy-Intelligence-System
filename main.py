import os
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

warnings.filterwarnings("ignore")


OUTPUT_DOCX = "BASIS_Final_10_Charts_Report.docx"

# Kaggle dataset
KAGGLE_DATASET = "surajjha101/bigbasket-entire-product-list-28k-datapoints"


# ============================================================
# 2. DOWNLOAD DATASET FROM KAGGLE
# ============================================================

try:
    import kagglehub
except ImportError:
    raise ImportError(
        "kagglehub is not installed.\n"
        "Install it first with:\n"
        "pip install kagglehub"
    )

print("Downloading BigBasket dataset from Kaggle...")

path = kagglehub.dataset_download(
    KAGGLE_DATASET
)

print("Kaggle dataset folder:")
print(path)


# ============================================================
# 3. FIND CSV FILE INSIDE KAGGLE DATASET
# ============================================================

dataset_path = Path(path)

csv_files = list(dataset_path.rglob("*.csv"))

if not csv_files:
    raise FileNotFoundError(
        "No CSV file was found inside the downloaded Kaggle dataset."
    )

# Prefer the main BigBasket CSV
bigbasket_csvs = [
    file for file in csv_files
    if "bigbasket" in file.name.lower()
]

if bigbasket_csvs:
    DATASET_PATH = bigbasket_csvs[0]
else:
    DATASET_PATH = csv_files[0]

print(f"Using CSV file:\n{DATASET_PATH}")


# ============================================================
# 4. LOAD DATA
# ============================================================

df = pd.read_csv(DATASET_PATH)

print(f"Original dataset shape: {df.shape}")

# ============================================================
# 5. CLEAN COLUMN NAMES
# ============================================================

df.columns = (
    df.columns
    .astype(str)
    .str.strip()
    .str.lower()
    .str.replace(" ", "_", regex=False)
)

# Remove unnecessary index column if present
if "index" in df.columns:
    df = df.drop(columns=["index"])


# ============================================================
# 6. CHECK REQUIRED COLUMNS
# ============================================================

required_cols = [
    "product",
    "category",
    "sub_category",
    "brand",
    "sale_price",
    "market_price",
    "rating",
]

missing_cols = [
    col for col in required_cols
    if col not in df.columns
]

if missing_cols:
    raise ValueError(
        f"\nMissing required columns: {missing_cols}\n"
        f"Available columns: {list(df.columns)}"
    )


# ============================================================
# 7. CLEAN DATA
# ============================================================

# Remove rows where important categorical fields are missing
df = df.dropna(
    subset=[
        "product",
        "brand",
        "category",
        "sub_category"
    ]
).copy()

# Convert numerical columns safely
df["sale_price"] = pd.to_numeric(
    df["sale_price"],
    errors="coerce"
)

df["market_price"] = pd.to_numeric(
    df["market_price"],
    errors="coerce"
)

df["rating"] = pd.to_numeric(
    df["rating"],
    errors="coerce"
)

# Remove rows with invalid prices
df = df.dropna(
    subset=[
        "sale_price",
        "market_price"
    ]
).copy()

df = df[
    (df["sale_price"] > 0) &
    (df["market_price"] > 0)
].copy()

# Ratings should be between 1 and 5
df.loc[
    (df["rating"] < 1) |
    (df["rating"] > 5),
    "rating"
] = np.nan

# Clean text fields
for col in [
    "product",
    "category",
    "sub_category",
    "brand"
]:
    df[col] = (
        df[col]
        .astype(str)
        .str.strip()
    )


# ============================================================
# 8. FEATURE ENGINEERING
# ============================================================

df["discount_amount"] = (
    df["market_price"] -
    df["sale_price"]
)

df["discount_percent"] = np.where(
    df["market_price"] > 0,
    (
        df["discount_amount"] /
        df["market_price"]
    ) * 100,
    0
)

df["discount_percent"] = (
    df["discount_percent"]
    .replace([np.inf, -np.inf], np.nan)
    .fillna(0)
)

# Avoid negative discount values
df["discount_percent"] = (
    df["discount_percent"]
    .clip(lower=0)
)


print(f"Cleaned dataset shape: {df.shape}")


# ============================================================
# 9. CREATE WORD DOCUMENT
# ============================================================

doc = Document()

section = doc.sections[0]

section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

# Default font
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10)


# ============================================================
# 10. COVER PAGE
# ============================================================

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = title.add_run("BASIS")
run.bold = True
run.font.size = Pt(28)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run(
    "BigBasket Assortment & Strategy Intelligence System"
)
run.bold = True
run.font.size = Pt(16)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = p.add_run(
    "10-Chart Product Assortment, Pricing, "
    "Brand and Category Analysis"
)
run.italic = True

doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER

info.add_run(
    f"Dataset records analysed: {len(df):,}\n"
    f"Categories: {df['category'].nunique():,} | "
    f"Subcategories: {df['sub_category'].nunique():,} | "
    f"Brands: {df['brand'].nunique():,}"
)

doc.add_paragraph(
    "Dataset scope note: This BigBasket dataset is a product "
    "catalogue snapshot. It does not contain customer IDs, "
    "orders, transaction dates, quantities sold, or actual "
    "revenue. Therefore, this report focuses on assortment, "
    "listed pricing, discount structure, ratings, brands and "
    "category positioning rather than claiming actual sales, "
    "revenue, demand or price elasticity."
)


# ============================================================
# 11. HELPER FUNCTIONS
# ============================================================

def add_heading(text, level=1):
    return doc.add_heading(text, level=level)


def add_description(
    purpose,
    look_for,
    business
):
    items = [
        ("Purpose", purpose),
        ("What to look for", look_for),
        ("Business relevance", business),
    ]

    for label, text in items:

        paragraph = doc.add_paragraph()

        run = paragraph.add_run(
            f"{label}: "
        )
        run.bold = True

        paragraph.add_run(text)


def add_chart(
    fig,
    title,
    chart_type,
    purpose,
    look_for,
    business
):

    # Temporary PNG file
    temp_file = (
        Path.cwd() /
        f"_basis_temp_chart_{len(doc.paragraphs)}.png"
    )

    # Save chart
    fig.savefig(
        temp_file,
        dpi=300,
        bbox_inches="tight"
    )

    # Close matplotlib figure
    plt.close(fig)

    # Chart heading
    add_heading(title, level=1)

    # Chart type
    paragraph = doc.add_paragraph()

    paragraph.add_run(
        "Chart type: "
    ).bold = True

    paragraph.add_run(chart_type)

    # Insert image into Word
    doc.add_picture(
        str(temp_file),
        width=Inches(6.8)
    )

    doc.paragraphs[-1].alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    # Description
    add_description(
        purpose,
        look_for,
        business
    )

    # Remove temporary image
    try:
        temp_file.unlink()
    except OSError:
        pass


# ============================================================
# CHART 1
# PRODUCT DISTRIBUTION ACROSS CATEGORIES
# ============================================================

category_counts = (
    df["category"]
    .value_counts()
    .sort_values()
)

fig, ax = plt.subplots(
    figsize=(10, 6)
)

ax.barh(
    category_counts.index,
    category_counts.values
)

ax.set_title(
    "Product Distribution Across Categories",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Number of Products"
)

ax.set_ylabel(
    "Category"
)

ax.grid(
    axis="x",
    alpha=0.25
)

fig.tight_layout()

add_chart(
    fig,
    "1. Product Distribution Across Categories",
    "Horizontal bar chart",
    "Shows how the catalogue is distributed across "
    "major product categories.",
    "Categories with very large or very small "
    "assortment sizes and the relative breadth "
    "of the catalogue.",
    "Useful for assortment planning and identifying "
    "categories with comparatively high or low "
    "catalogue coverage."
)


# ============================================================
# CHART 2
# TOP 15 SUBCATEGORIES
# ============================================================

subcat_counts = (
    df["sub_category"]
    .value_counts()
    .head(15)
    .sort_values()
)

fig, ax = plt.subplots(
    figsize=(10, 7)
)

ax.barh(
    subcat_counts.index,
    subcat_counts.values
)

ax.set_title(
    "Top 15 Subcategories by Product Count",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Number of Products"
)

ax.set_ylabel(
    "Subcategory"
)

ax.grid(
    axis="x",
    alpha=0.25
)

fig.tight_layout()

add_chart(
    fig,
    "2. Top 15 Subcategories by Product Count",
    "Horizontal bar chart",
    "Highlights the subcategories with the largest "
    "number of listed products.",
    "Which subcategories dominate the assortment "
    "and how concentrated the catalogue is among "
    "leading subcategories.",
    "Supports assortment depth decisions and helps "
    "identify areas where catalogue breadth is "
    "already high."
)


# ============================================================
# CHART 3
# SALE PRICE DISTRIBUTION
# ============================================================

fig, ax = plt.subplots(
    figsize=(10, 6)
)

ax.hist(
    df["sale_price"],
    bins=60,
    edgecolor="black",
    alpha=0.75
)

ax.set_title(
    "Distribution of Product Sale Prices",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Sale Price (INR)"
)

ax.set_ylabel(
    "Number of Products"
)

ax.grid(
    axis="y",
    alpha=0.2
)

# Avoid extreme high-price values compressing the chart
q99 = df["sale_price"].quantile(0.99)

if pd.notna(q99) and q99 > 0:
    ax.set_xlim(
        0,
        q99
    )

fig.tight_layout()

add_chart(
    fig,
    "3. Distribution of Product Sale Prices",
    "Histogram",
    "Shows the overall spread of listed selling "
    "prices in the product catalogue.",
    "The concentration of products in lower, middle "
    "and higher price ranges and the degree of "
    "price skew.",
    "Useful for understanding the catalogue's price "
    "architecture and identifying dominant price bands."
)


# ============================================================
# CHART 4
# SALE PRICE DISTRIBUTION ACROSS CATEGORIES
# ============================================================

plot_df = df.copy()

upper = plot_df["sale_price"].quantile(0.98)

plot_df["sale_price_plot"] = (
    plot_df["sale_price"]
    .clip(upper=upper)
)

order = (
    plot_df
    .groupby("category")["sale_price"]
    .median()
    .sort_values()
    .index
    .tolist()
)

# Create clean boxplot data
box_data = []

valid_labels = []

for category in order:

    values = (
        plot_df.loc[
            plot_df["category"] == category,
            "sale_price_plot"
        ]
        .dropna()
        .values
    )

    if len(values) > 0:
        box_data.append(values)
        valid_labels.append(category)


fig, ax = plt.subplots(
    figsize=(11, 7)
)

# IMPORTANT:
# Use tick_labels instead of labels.
# "labels" is deprecated in newer Matplotlib versions.
ax.boxplot(
    box_data,
    vert=False,
    tick_labels=valid_labels,
    showfliers=False
)

ax.set_title(
    "Sale Price Distribution Across Categories",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Sale Price (INR)"
)

ax.set_ylabel(
    "Category"
)

ax.grid(
    axis="x",
    alpha=0.25
)

fig.tight_layout()

add_chart(
    fig,
    "4. Sale Price Distribution Across Categories",
    "Horizontal boxplot",
    "Compares the distribution and central tendency "
    "of listed sale prices across categories.",
    "Differences in median prices, spread and "
    "category-level price ranges.",
    "Helps compare category price positioning and "
    "identify categories with broader or narrower "
    "price structures."
)


# ============================================================
# CHART 5
# MARKET PRICE VS SALE PRICE
# ============================================================

sample_size = min(
    len(df),
    8000
)

scatter_df = df.sample(
    sample_size,
    random_state=42
)

fig, ax = plt.subplots(
    figsize=(9, 7)
)

ax.scatter(
    scatter_df["market_price"],
    scatter_df["sale_price"],
    alpha=0.35,
    s=18
)

combined_prices = np.concatenate(
    [
        scatter_df["market_price"].values,
        scatter_df["sale_price"].values
    ]
)

max_price = np.nanpercentile(
    combined_prices,
    98
)

if max_price > 0:

    ax.plot(
        [0, max_price],
        [0, max_price],
        linestyle="--",
        linewidth=1.5
    )

    ax.set_xlim(
        0,
        max_price
    )

    ax.set_ylim(
        0,
        max_price
    )

ax.set_title(
    "Market Price vs Sale Price",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Market Price (INR)"
)

ax.set_ylabel(
    "Sale Price (INR)"
)

ax.grid(
    alpha=0.2
)

fig.tight_layout()

add_chart(
    fig,
    "5. Market Price vs Sale Price",
    "Scatter plot",
    "Examines the relationship between listed market "
    "price and listed sale price.",
    "How closely sale prices track market prices "
    "and where larger price gaps appear.",
    "Useful for reviewing pricing and discount "
    "structure across the catalogue. It should not "
    "be interpreted as actual realised sales or revenue."
)


# ============================================================
# CHART 6
# AVERAGE DISCOUNT BY CATEGORY
# ============================================================

category_discount = (
    df.groupby("category")["discount_percent"]
    .mean()
    .sort_values()
)

fig, ax = plt.subplots(
    figsize=(10, 6)
)

ax.barh(
    category_discount.index,
    category_discount.values
)

ax.set_title(
    "Average Discount Percentage by Category",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Average Discount (%)"
)

ax.set_ylabel(
    "Category"
)

ax.grid(
    axis="x",
    alpha=0.25
)

fig.tight_layout()

add_chart(
    fig,
    "6. Average Discount Percentage by Category",
    "Horizontal bar chart",
    "Compares the average listed discount percentage "
    "across categories.",
    "Categories with relatively higher or lower "
    "discount intensity.",
    "Supports category-level pricing and promotional "
    "strategy analysis using catalogue price information."
)


# ============================================================
# CHART 7
# PRODUCT RATING DISTRIBUTION
# ============================================================

rating_df = (
    df.dropna(
        subset=["rating"]
    )
)

fig, ax = plt.subplots(
    figsize=(9, 6)
)

bins = np.arange(
    0.75,
    5.26,
    0.25
)

ax.hist(
    rating_df["rating"],
    bins=bins,
    edgecolor="black",
    alpha=0.75
)

ax.set_title(
    "Distribution of Product Ratings",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Customer Rating"
)

ax.set_ylabel(
    "Number of Products"
)

ax.set_xticks(
    np.arange(
        1,
        5.1,
        0.5
    )
)

ax.grid(
    axis="y",
    alpha=0.2
)

fig.tight_layout()

add_chart(
    fig,
    "7. Distribution of Product Ratings",
    "Histogram",
    "Shows how product ratings are distributed "
    "across the catalogue.",
    "Whether ratings are concentrated toward higher "
    "or lower values and how widely ratings vary.",
    "Provides a catalogue-level view of product "
    "rating patterns that can support assortment "
    "quality review."
)


# ============================================================
# CHART 8
# TOP 15 BRANDS BY PRODUCT LISTINGS
# ============================================================

brand_counts = (
    df["brand"]
    .value_counts()
    .head(15)
    .sort_values()
)

fig, ax = plt.subplots(
    figsize=(10, 7)
)

ax.barh(
    brand_counts.index,
    brand_counts.values
)

ax.set_title(
    "Top 15 Brands by Product Listings",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Number of Listed Products"
)

ax.set_ylabel(
    "Brand"
)

ax.grid(
    axis="x",
    alpha=0.25
)

fig.tight_layout()

add_chart(
    fig,
    "8. Top 15 Brands by Product Listings",
    "Horizontal bar chart",
    "Identifies brands with the largest number "
    "of listed products in the catalogue.",
    "Brands with the broadest assortment footprint.",
    "Useful for understanding brand assortment "
    "concentration and catalogue representation."
)


# ============================================================
# CHART 9
# BRAND POSITIONING
# ============================================================

brand_position = (
    df.groupby("brand")
    .agg(
        avg_price=("sale_price", "mean"),
        avg_rating=("rating", "mean"),
        product_count=("product", "count")
    )
    .reset_index()
)

# Only brands with at least 20 products
brand_position = brand_position[
    brand_position["product_count"] >= 20
].copy()

brand_position = brand_position.dropna(
    subset=["avg_rating"]
)

fig, ax = plt.subplots(
    figsize=(11, 8)
)

if not brand_position.empty:

    sizes = np.clip(
        brand_position["product_count"] * 2.5,
        30,
        500
    )

    ax.scatter(
        brand_position["avg_price"],
        brand_position["avg_rating"],
        s=sizes,
        alpha=0.65
    )

    # Label top 15 brands by listing count
    labels = (
        brand_position
        .nlargest(
            min(15, len(brand_position)),
            "product_count"
        )
    )

    for _, row in labels.iterrows():

        ax.annotate(
            str(row["brand"]),
            (
                row["avg_price"],
                row["avg_rating"]
            ),
            xytext=(5, 5),
            textcoords="offset points",
            fontsize=8
        )

    min_rating = brand_position[
        "avg_rating"
    ].min()

    max_rating = brand_position[
        "avg_rating"
    ].max()

    lower_limit = max(
        1,
        min_rating - 0.1
    )

    upper_limit = min(
        5.05,
        max_rating + 0.1
    )

    if lower_limit < upper_limit:
        ax.set_ylim(
            lower_limit,
            upper_limit
        )

else:
    ax.text(
        0.5,
        0.5,
        "No brands with 20 or more products",
        ha="center",
        va="center",
        transform=ax.transAxes
    )

ax.set_title(
    "Brand Positioning: Average Price vs Average Rating",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Average Sale Price (INR)"
)

ax.set_ylabel(
    "Average Rating"
)

ax.grid(
    alpha=0.2
)

fig.tight_layout()

add_chart(
    fig,
    "9. Brand Positioning: Average Price vs Average Rating",
    "Scatter plot with point size representing product listings",
    "Positions sufficiently represented brands using "
    "average listed price and average product rating.",
    "Brands that combine different price levels and "
    "rating levels, plus the relative size of their "
    "catalogue presence.",
    "Supports brand positioning analysis without "
    "claiming actual brand sales performance."
)


# ============================================================
# CHART 10
# CATEGORY POSITIONING
# ============================================================

category_position = (
    df.groupby("category")
    .agg(
        avg_price=("sale_price", "mean"),
        avg_discount=("discount_percent", "mean"),
        product_count=("product", "count")
    )
    .reset_index()
)

fig, ax = plt.subplots(
    figsize=(11, 8)
)

sizes = np.clip(
    category_position["product_count"] * 1.2,
    60,
    900
)

ax.scatter(
    category_position["avg_price"],
    category_position["avg_discount"],
    s=sizes,
    alpha=0.7
)

for _, row in category_position.iterrows():

    ax.annotate(
        str(row["category"]),
        (
            row["avg_price"],
            row["avg_discount"]
        ),
        xytext=(5, 5),
        textcoords="offset points",
        fontsize=8
    )

ax.set_title(
    "Category Positioning: Price vs Discount vs Assortment",
    fontsize=15,
    fontweight="bold"
)

ax.set_xlabel(
    "Average Sale Price (INR)"
)

ax.set_ylabel(
    "Average Discount (%)"
)

ax.grid(
    alpha=0.2
)

fig.tight_layout()

add_chart(
    fig,
    "10. Category Positioning: Price vs Discount vs Assortment",
    "Scatter plot with point size representing product listings",
    "Compares categories using average listed sale price, "
    "average discount and assortment size.",
    "Categories with different combinations of price "
    "level, discount intensity and catalogue breadth.",
    "Provides a compact view for category assortment "
    "and pricing strategy discussions."
)


# ============================================================
# FINAL SUMMARY
# ============================================================

add_heading(
    "Final Summary",
    level=1
)

summary = doc.add_paragraph()

summary.add_run(
    "The ten charts provide a catalogue-level view of "
    "BigBasket's product assortment, pricing structure, "
    "discount patterns, ratings, brand representation "
    "and category positioning. The analysis is based "
    "on listed product information rather than "
    "transactional sales data."
)


# ============================================================
# SUMMARY TABLE
# ============================================================

table = doc.add_table(
    rows=1,
    cols=3
)

table.style = "Table Grid"

header = table.rows[0].cells

header[0].text = "Chart"
header[1].text = "Focus"
header[2].text = "Primary Use"

summary_rows = [
    (
        "1",
        "Category assortment",
        "Assortment breadth"
    ),
    (
        "2",
        "Subcategory assortment",
        "Assortment depth"
    ),
    (
        "3",
        "Sale price distribution",
        "Price architecture"
    ),
    (
        "4",
        "Category price distribution",
        "Category price positioning"
    ),
    (
        "5",
        "Market vs sale price",
        "Pricing / discount structure"
    ),
    (
        "6",
        "Category discount",
        "Promotional intensity"
    ),
    (
        "7",
        "Product ratings",
        "Catalogue rating pattern"
    ),
    (
        "8",
        "Brand listings",
        "Brand assortment footprint"
    ),
    (
        "9",
        "Brand positioning",
        "Price vs rating positioning"
    ),
    (
        "10",
        "Category positioning",
        "Price vs discount vs assortment"
    ),
]

for row in summary_rows:

    cells = table.add_row().cells

    for i, value in enumerate(row):
        cells[i].text = value


# ============================================================
# IMPORTANT DATASET NOTE
# ============================================================

doc.add_paragraph()

note = doc.add_paragraph()

run = note.add_run(
    "Important: This dataset does not contain transaction "
    "quantity, order history, customer IDs or transaction "
    "dates. Listed product counts, prices, discounts and "
    "ratings should therefore be interpreted as catalogue "
    "indicators, not actual sales or revenue measures."
)

run.italic = True


# ============================================================
# SAVE DOCX
# ============================================================

doc.save(
    OUTPUT_DOCX
)

print("\n" + "=" * 70)
print("SUCCESS")
print("=" * 70)
print(f"Word report created: {OUTPUT_DOCX}")
print(f"Dataset used: {DATASET_PATH}")
print("Exactly 10 charts were inserted into the Word document.")
print("=" * 70)
